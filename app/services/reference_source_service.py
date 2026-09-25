"""
ReferenceSourceService — Multi-Modal Reference Source Parser
Handles:
1. Jira Attachments: .md/.txt (direct text), .docx (paragraph/table extraction),
   Images (.png/.jpg/.jpeg/.webp) — described via Gemini Vision (LLMService)
2. Git Repository Links: fetches referenced files/specs from GitHub using TokenSecretService
3. Knowledge Base Integration: indexes extracted reference texts into ChromaDB
"""
import os
import re
import logging
import requests

logger = logging.getLogger(__name__)


class ReferenceSourceService:

    @classmethod
    def fetch_and_index_references(cls, story, workflow_id: int = None) -> dict:
        """
        Main entry point. Aggregates all reference sources for a story and
        indexes them into ChromaDB collection `story_refs_{story_id}`.

        Returns:
        {
            "sources": [{"type": str, "filename": str, "text": str}, ...],
            "collection_name": "story_refs_<story_id>",
            "indexed_count": int
        }
        """
        story_id = story.id if hasattr(story, 'id') else story.get('id', 0)
        description = story.description if hasattr(story, 'description') else story.get('description', '')
        acceptance_criteria = story.acceptance_criteria if hasattr(story, 'acceptance_criteria') else story.get('acceptance_criteria', '')
        jira_key = story.jira_story_key if hasattr(story, 'jira_story_key') else story.get('jira_story_key', '')

        sources = []

        # 1. Parse Jira attachments
        if jira_key:
            jira_sources = cls._fetch_jira_attachments(jira_key)
            sources.extend(jira_sources)

        # 2. Detect and fetch Git links from description + AC
        combined_text = f"{description or ''}\n{acceptance_criteria or ''}"
        git_sources = cls._fetch_git_links(combined_text)
        sources.extend(git_sources)

        # 3. Detect and fetch Reference Repository files if provided
        ref_sources = cls._fetch_reference_repo_sources(story, combined_text, workflow_id)
        sources.extend(ref_sources)

        # 4. Index all extracted references into ChromaDB
        collection_name = f"story_refs_{story_id}"
        indexed_count = cls._index_to_chromadb(collection_name, sources)

        logger.info(
            f"[ReferenceSourceService] Story {story_id}: indexed {indexed_count} reference "
            f"chunks from {len(sources)} source(s) into '{collection_name}'"
        )

        return {
            "sources": sources,
            "collection_name": collection_name,
            "indexed_count": indexed_count
        }

    # ─────────────────────────────────────────────────────────────────────────
    # Jira Attachments
    # ─────────────────────────────────────────────────────────────────────────

    @classmethod
    def _fetch_jira_attachments(cls, jira_key: str) -> list:
        """
        Fetch all attachments from a Jira issue and parse each by type.
        Returns list of {type, filename, text} dicts.
        """
        sources = []
        try:
            from app.services.jira_service import JiraService
            issue = JiraService.get_issue(jira_key)
            if not issue:
                return sources

            attachments = (issue.get('fields') or {}).get('attachment') or []
            for att in attachments:
                filename = att.get('filename', '')
                content_url = att.get('content', '')
                mime_type = (att.get('mimeType') or '').lower()

                if not content_url or not filename:
                    continue

                ext = os.path.splitext(filename)[1].lower()
                try:
                    raw = cls._download_attachment(content_url)
                    if raw is None:
                        continue

                    if ext in ('.md', '.txt'):
                        text = cls._parse_text(raw)
                        sources.append({"type": "markdown", "filename": filename, "text": text})

                    elif ext in ('.docx',):
                        text = cls._parse_docx(raw)
                        sources.append({"type": "docx", "filename": filename, "text": text})

                    elif ext in ('.png', '.jpg', '.jpeg', '.webp', '.gif'):
                        text = cls._describe_image_via_vision(raw, mime_type or "image/png", filename)
                        sources.append({"type": "image", "filename": filename, "text": text})

                    else:
                        logger.debug(f"[ReferenceSourceService] Skipping unsupported attachment type: {filename}")

                except Exception as att_err:
                    logger.warning(f"[ReferenceSourceService] Failed to process attachment '{filename}': {att_err}")

        except Exception as e:
            logger.warning(f"[ReferenceSourceService] Failed to fetch Jira attachments for {jira_key}: {e}")

        return sources

    @classmethod
    def _download_attachment(cls, url: str) -> bytes:
        """Download an attachment from Jira using the configured auth credentials."""
        try:
            from flask import current_app
            from requests.auth import HTTPBasicAuth
            auth = HTTPBasicAuth(
                current_app.config.get('JIRA_EMAIL', ''),
                current_app.config.get('JIRA_API_TOKEN', '')
            )
            resp = requests.get(url, auth=auth, timeout=30)
            if resp.status_code == 200:
                return resp.content
            logger.warning(f"[ReferenceSourceService] Attachment download returned {resp.status_code}: {url}")
        except Exception as e:
            logger.warning(f"[ReferenceSourceService] Attachment download failed: {e}")
        return None

    @staticmethod
    def _parse_text(raw: bytes) -> str:
        """Parse raw bytes as UTF-8 text (for .md / .txt files)."""
        for encoding in ('utf-8', 'utf-8-sig', 'latin-1'):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode('utf-8', errors='replace')

    @staticmethod
    def _parse_docx(raw: bytes) -> str:
        """
        Extract paragraphs and tables from a .docx file.
        Uses python-docx if available, falls back to raw XML extraction.
        """
        try:
            import io
            from docx import Document
            doc = Document(io.BytesIO(raw))
            parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            return "\n".join(parts)
        except ImportError:
            # Fallback: raw XML text extraction (no python-docx)
            import io
            import zipfile
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                    with z.open('word/document.xml') as f:
                        xml = f.read().decode('utf-8', errors='replace')
                text = re.sub(r'<[^>]+>', ' ', xml)
                text = re.sub(r'\s+', ' ', text).strip()
                return text
            except Exception:
                return "[DOCX parsing failed — python-docx not installed]"
        except Exception as e:
            logger.warning(f"[ReferenceSourceService] DOCX parse error: {e}")
            return ""

    @staticmethod
    def _describe_image_via_vision(raw: bytes, mime_type: str, filename: str) -> str:
        """
        Use Gemini Vision (LLMService) to generate a detailed technical description
        of a UI mockup, architecture diagram, or flowchart image.
        """
        try:
            from app.services.llm_service import LLMService
            prompt = (
                f"You are analyzing a technical image attachment named '{filename}' from a Jira story. "
                "Describe in detail what this image shows, including: UI components, layout, data flows, "
                "architecture layers, API interactions, or business logic visible in the diagram. "
                "Be specific and technical — your description will be used as context for a software developer "
                "implementing the feature shown in this image."
            )
            description = LLMService.generate_with_image(
                prompt=prompt,
                image_data=raw,
                mime_type=mime_type,
                agent_name="ReferenceSourceService"
            )
            return description or f"[Image: {filename} — vision description unavailable]"
        except Exception as e:
            logger.warning(f"[ReferenceSourceService] Vision description failed for '{filename}': {e}")
            return f"[Image: {filename} — vision description failed: {e}]"

    # ─────────────────────────────────────────────────────────────────────────
    # Git Repository Links
    # ─────────────────────────────────────────────────────────────────────────

    @classmethod
    def _fetch_git_links(cls, text: str) -> list:
        """
        Detect GitHub file/blob links in text and fetch their content.
        Returns list of {type, filename, text} dicts.
        """
        sources = []
        if not text:
            return sources

        # Detect GitHub blob/raw file URLs
        github_pattern = re.compile(
            r'https://github\.com/([^/\s]+)/([^/\s]+)/(?:blob|raw)/([^/\s]+)/([^\s\)>\"\']+)',
            re.IGNORECASE
        )

        seen_urls = set()
        for match in github_pattern.finditer(text):
            org, repo, branch, filepath = match.group(1), match.group(2), match.group(3), match.group(4)
            raw_url = f"https://raw.githubusercontent.com/{org}/{repo}/{branch}/{filepath}"

            if raw_url in seen_urls:
                continue
            seen_urls.add(raw_url)

            try:
                from app.services.token_secret_service import TokenSecretService
                repo_url = f"https://github.com/{org}/{repo}"
                token = TokenSecretService.get_token_for_repo(repo_url)
                headers = {}
                if token:
                    headers["Authorization"] = f"Bearer {token}"

                resp = requests.get(raw_url, headers=headers, timeout=15)
                if resp.status_code == 200:
                    content = resp.text
                    sources.append({
                        "type": "git_link",
                        "filename": f"{org}/{repo}/{filepath}",
                        "text": content[:10000]  # Limit to 10K chars
                    })
                    logger.info(f"[ReferenceSourceService] Fetched git link: {raw_url}")
                else:
                    logger.warning(f"[ReferenceSourceService] Git link fetch returned {resp.status_code}: {raw_url}")
            except Exception as e:
                logger.warning(f"[ReferenceSourceService] Failed to fetch git link {raw_url}: {e}")

        return sources

    # ─────────────────────────────────────────────────────────────────────────
    # Reference Repository Ingestion
    # ─────────────────────────────────────────────────────────────────────────

    @classmethod
    def _fetch_reference_repo_sources(cls, story, text: str, workflow_id: int = 0) -> list:
        sources = []
        try:
            from app.services.token_secret_service import TokenSecretService
            from flask import current_app
            import subprocess

            ref_name = None
            ref_url = None
            ref_branch = 'main'

            # 1. Check story.repository_details
            details = getattr(story, 'repository_details', None) if story else None
            if isinstance(details, list):
                for d in details:
                    if isinstance(d, dict) and d.get('is_reference'):
                        ref_name = d.get('name')
                        ref_url = d.get('url')
                        ref_branch = d.get('branch', 'main')
                        break

            # 2. Check text for Reference Repo patterns
            if not ref_url and text:
                m = re.search(r'(?:Reference\s*Repo(?:sitory)?|Ref\s*Repo(?:sitory)?)\s*[:\-]\s*([^\s\n\r]+)', text, re.IGNORECASE)
                if m:
                    val = m.group(1).strip().strip('`').strip('"').strip("'")
                    cfg = TokenSecretService.get_repo_config(val)
                    if cfg:
                        ref_name = val
                        ref_url = cfg.get('url')
                        ref_branch = cfg.get('branch', 'main')
                    elif val.startswith('http') or val.startswith('git@'):
                        ref_url = val
                        ref_name = val.split('/')[-1].replace('.git', '')

            if not ref_url:
                return sources

            token = TokenSecretService.get_token_for_repo(ref_url) or current_app.config.get('GITHUB_TOKEN')
            clone_url = ref_url
            if token and "github.com" in ref_url:
                clone_url = ref_url.replace("https://github.com/", f"https://x-access-token:{token}@github.com/")

            from app.config.settings import Config
            workspaces_root = getattr(Config, 'WORKSPACES_DIR', current_app.config.get('WORKSPACES_DIR', os.path.abspath(os.path.join(current_app.root_path, '..', 'workspaces'))))
            os.makedirs(workspaces_root, exist_ok=True)
            ref_dir = os.path.join(workspaces_root, f"ref_repo_{workflow_id}_{ref_name or 'ref'}")

            git_env = dict(os.environ)
            git_env['GIT_TERMINAL_PROMPT'] = '0'

            if not os.path.exists(ref_dir):
                cmd = ['git', 'clone', '--depth', '20', '-b', ref_branch, clone_url, ref_dir]
                proc = subprocess.run(cmd, capture_output=True, text=True, timeout=60, env=git_env)
                if proc.returncode != 0:
                    logger.warning(f"[ReferenceSourceService] Failed to clone ref branch '{ref_branch}': {proc.stderr}. Retrying default branch clone...")
                    cmd_fb = ['git', 'clone', '--depth', '20', clone_url, ref_dir]
                    proc_fb = subprocess.run(cmd_fb, capture_output=True, text=True, timeout=60, env=git_env)
                    if proc_fb.returncode != 0:
                        logger.error(f"[ReferenceSourceService] Fallback clone also failed: {proc_fb.stderr}")

            if os.path.exists(ref_dir):
                for root, _, files in os.walk(ref_dir):
                    if any(ig in root for ig in ('.git', 'node_modules', 'venv', '__pycache__', 'dist', 'build', '.idea', '.vscode')):
                        continue
                    for f in files:
                        if f.endswith(('.py', '.js', '.ts', '.jsx', '.tsx', '.json', '.md', '.sql', '.html')):
                            fp = os.path.join(root, f)
                            rel = os.path.relpath(fp, ref_dir).replace('\\', '/')
                            try:
                                with open(fp, 'r', encoding='utf-8', errors='replace') as r_file:
                                    cnt = r_file.read()
                                if cnt.strip() and len(cnt) < 15000:
                                    sources.append({
                                        "type": "reference_repo",
                                        "filename": f"ref:{ref_name}/{rel}",
                                        "text": cnt
                                    })
                            except Exception:
                                pass
                logger.info(f"[ReferenceSourceService] Indexed {len(sources)} files from reference repository '{ref_name}'")
        except Exception as e:
            logger.warning(f"[ReferenceSourceService] Reference repo fetch failed (non-fatal): {e}")

        return sources

    # ─────────────────────────────────────────────────────────────────────────
    # ChromaDB Knowledge Base Indexing
    # ─────────────────────────────────────────────────────────────────────────

    @classmethod
    def _index_to_chromadb(cls, collection_name: str, sources: list) -> int:
        """
        Index all extracted reference texts into a ChromaDB collection.
        Returns the number of chunks indexed.
        """
        if not sources:
            return 0

        try:
            from app.services.rag_service import RagService
            rag = RagService.get_instance()

            # Get or create the story refs collection
            try:
                collection = rag.client.get_collection(collection_name)
            except Exception:
                collection = rag.client.create_collection(
                    name=collection_name,
                    metadata={"hnsw:space": "cosine"}
                )

            documents = []
            metadatas = []
            ids = []

            for i, source in enumerate(sources):
                text = (source.get('text') or '').strip()
                if not text:
                    continue

                # Chunk large texts
                chunk_size = 2000
                chunks = [text[j:j + chunk_size] for j in range(0, len(text), chunk_size)]
                for k, chunk in enumerate(chunks):
                    if not chunk.strip():
                        continue
                    documents.append(chunk)
                    metadatas.append({
                        "source_type": source.get("type", "unknown"),
                        "filename": source.get("filename", ""),
                        "chunk_index": k
                    })
                    ids.append(f"{collection_name}_src{i}_chunk{k}")

            if documents:
                collection.add(documents=documents, metadatas=metadatas, ids=ids)
                logger.info(
                    f"[ReferenceSourceService] Indexed {len(documents)} chunks into '{collection_name}'"
                )
            return len(documents)

        except Exception as e:
            logger.error(f"[ReferenceSourceService] ChromaDB indexing failed: {e}")
            return 0

    @classmethod
    def query_references(cls, story_id: int, query: str, n_results: int = 5) -> list:
        """
        Query the story reference collection for relevant context.
        Returns list of {text, source_type, filename} dicts.
        """
        collection_name = f"story_refs_{story_id}"
        try:
            from app.services.rag_service import RagService
            rag = RagService.get_instance()
            try:
                collection = rag.client.get_collection(collection_name)
            except Exception:
                return []

            results = collection.query(query_texts=[query], n_results=min(n_results, 10))
            docs = results.get('documents', [[]])[0]
            metas = results.get('metadatas', [[]])[0]

            return [
                {
                    "text": doc,
                    "source_type": meta.get("source_type", "unknown"),
                    "filename": meta.get("filename", "")
                }
                for doc, meta in zip(docs, metas)
            ]
        except Exception as e:
            logger.warning(f"[ReferenceSourceService] Query failed for '{collection_name}': {e}")
            return []

